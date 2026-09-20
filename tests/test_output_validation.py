import csv
import tempfile
import unittest
from datetime import date, timedelta
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

from models import BertMember, HourCaps, Schedule, Volunteer
from output import collect_warnings, export_master_schedule_csv, export_schedule_xlsx
from validation import validate_schedule

D = date(2026, 9, 21)


class ValidationTests(unittest.TestCase):
    def test_all_daytime_subsets_against_explicit_allowed_combinations(self):
        allowed = {frozenset(s.split()) for s in ['', 'A', 'B', 'C', 'D', 'AM', 'PM',
                    'A B', 'B C', 'C D', 'AM C', 'B PM', 'AM PM',
                    'A B C', 'B C D', 'AM C D', 'A B PM', 'A B C D']}
        kinds = ('A', 'B', 'C', 'D', 'AM', 'PM')
        for count in range(7):
            for subset in combinations(kinds, count):
                with self.subTest(subset=subset):
                    p = Volunteer('Test', 'Person', 'test@example.com')
                    p.assigned = [(D, k) for k in subset if k in ('AM', 'PM')]
                    p.campus_assigned = [(D, k) for k in subset if k not in ('AM', 'PM')]
                    p.available = set(p.assigned)
                    p.campus_available = set(p.campus_assigned)
                    schedule = Schedule({k: [p] for k in p.assigned}, {k: [p] for k in p.campus_assigned})
                    errors = validate_schedule(schedule, [p], {k: 'BLS' for k in p.assigned},
                                               p.campus_assigned, HourCaps(18, 12, 9), 2)
                    self.assertEqual(not errors, frozenset(subset) in allowed)

    def test_checker_catches_night_and_next_day_campus(self):
        p = Volunteer('Test', 'Person', 'test@example.com')
        night, campus = (D, 'NIGHT'), (D + timedelta(days=1), 'D')
        p.assigned, p.campus_assigned = [night], [campus]
        p.available, p.campus_available = {night}, {campus}
        errors = validate_schedule(Schedule({night: [p]}, {campus: [p]}), [p],
                                   {night: 'BLS'}, [campus], HourCaps(), 2)
        self.assertTrue(any('rest' in e for e in errors))

    def test_checker_catches_duplicated_seat_and_unsynced_totals(self):
        p = Volunteer('Test', 'Person', 'test@example.com', available={(D, 'AM')})
        errors = validate_schedule(Schedule({(D, 'AM'): [p, p]}, {}), [p],
                                   {(D, 'AM'): 'BLS'}, [], HourCaps(), 2)
        self.assertTrue(any('multiple seats' in e for e in errors))
        self.assertTrue(any('totals' in e for e in errors))


class OutputTests(unittest.TestCase):
    def export(self, schedule, providers, capacity=2):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'master.csv'
            export_master_schedule_csv(schedule, providers, D, path, 'F26B2', 921, capacity)
            with path.open(newline='') as f:
                return list(csv.DictReader(f))

    def test_weekend_bls_auth_is_r1_crew_with_no_utility_vehicle(self):
        key = (D + timedelta(days=4), 'NIGHT')
        auth = Volunteer('Auth', 'Person', 'auth@example.com', 'Auth')
        rows = self.export(Schedule({key: [auth]}, {}), {key: 'BLS'})
        self.assertEqual([(r['Vehicle'], r['Seat'], r['Requires'], r['Assigned/Name']) for r in rows],
                         [('R1', 'C2', 'CREW', 'Auth Person'),
                          ('R1', 'C3', 'CREW', ''), ('R1', 'C4', 'CREW', '')])

    def test_als_keeps_evdt_truck_seat_open_when_only_auth_available(self):
        key = (D, 'AM')
        auth = Volunteer('Auth', 'Person', 'auth@example.com', 'Auth')
        rows = self.export(Schedule({key: [auth]}, {}), {key: 'ALS'})
        self.assertEqual((rows[0]['Vehicle'], rows[0]['Requires'], rows[0]['Assigned/Name']), ('R1', 'EVDT', ''))
        self.assertEqual(rows[1]['Assigned/Name'], 'Auth Person')

    def test_weekday_bls_crew_uses_the_v3_visible_c2_row(self):
        key = (D, 'AM')
        person = Volunteer('One', 'EMT', 'one@example.com')
        rows = self.export(Schedule({key: [person]}, {}), {key: 'BLS'})
        self.assertEqual([(r['Vehicle'], r['Seat'], r['Requires'], r['Assigned/Name']) for r in rows],
                         [('R1', 'C2', 'CREW', 'One EMT')])

    def test_als_weekend_keeps_evdt_driver_and_numbers_crew_from_c2(self):
        key = (D + timedelta(days=4), 'NIGHT')
        evdt = Volunteer('EVDT', 'Person', 'evdt@example.com', 'EVDT')
        auth = Volunteer('Auth', 'Person', 'auth@example.com', 'Auth')
        rows = self.export(Schedule({key: [evdt, auth]}, {}), {key: 'ALS'})
        assigned = {r['Assigned/Name']: r for r in rows if r['Assigned/Name']}
        self.assertEqual((assigned['EVDT Person']['Vehicle'], assigned['EVDT Person']['Seat']), ('R1', 'Driver'))
        self.assertEqual((assigned['Auth Person']['Vehicle'], assigned['Auth Person']['Seat']), ('R1', 'C2'))
        self.assertTrue(all(r['Vehicle'] == 'R1' for r in rows))

    def test_export_respects_campus_capacity_and_keeps_everyone(self):
        people = [BertMember(str(i), 'Campus', f'{i}@example.com') for i in range(3)]
        rows = self.export(Schedule({}, {(D, 'A'): people}), {}, 3)
        self.assertEqual({r['Assigned/Name'] for r in rows}, {p.full_name for p in people})
        self.assertEqual(len({r['ShiftID'] for r in rows}), 3)

    def test_export_refuses_to_drop_people(self):
        people = [BertMember(str(i), 'Campus', f'{i}@example.com') for i in range(3)]
        with self.assertRaisesRegex(ValueError, 'drop'):
            self.export(Schedule({}, {(D, 'A'): people}), {}, 2)

    def test_no_false_driver_or_second_campus_responder_warning(self):
        p = Volunteer('Test', 'Person', 'test@example.com')
        b = BertMember('Campus', 'Only', 'campus@example.com')
        issues = collect_warnings(Schedule({(D, 'AM'): [p]}, {(D, 'A'): [b]}),
                                  {(D, 'AM'): 'BLS'}, [], HourCaps())
        self.assertEqual(issues, [])

    def test_campus_and_ambulance_shortfalls_are_reported_separately(self):
        p = Volunteer('Test', 'Person', 'test@example.com')
        issues = collect_warnings(Schedule({}, {}), {}, [p], HourCaps())
        self.assertEqual({r[0] for r in issues}, {'AMBULANCE UNDER HOURS', 'CAMPUS UNDER HOURS'})

    def test_workbook_contains_all_views_and_no_strike_list(self):
        p = Volunteer('Test', 'Person', 'test@example.com')
        with tempfile.TemporaryDirectory() as tmp:
            path = export_schedule_xlsx(Schedule({(D, 'AM'): [p]}, {(D, 'C'): []}),
                                        [p], {(D, 'AM'): 'BLS'}, HourCaps(), Path(tmp) / 'schedule.xlsx')
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ['Schedule', 'Campus Response', 'Hour Summary', 'Warnings', 'Solver'])
            self.assertEqual(wb['Hour Summary']['I2'].value, 18)
            self.assertEqual(wb['Hour Summary']['J2'].value, 6)
            wb.close()


if __name__ == '__main__':
    unittest.main()
